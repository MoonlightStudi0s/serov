#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка отчёта по лабораторной работе № 2 в ODT, DOCX и PDF.

Первичным форматом является ODT. Для него используется styles.xml и макет
страницы из lab-01/examples/Лаба 2.odt: A4, поля 30/15/20/20 мм, Times New
Roman 14 пт, полуторный интервал, абзац 1,25 см и номера страниц внизу по
центру (титульный лист без номера). Это сохраняет оформление исходного
шаблона, а не только приблизительно повторяет его.

Зависимости: python-docx, reportlab, Pillow (см. requirements.txt).
"""

from __future__ import annotations

import copy
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image

BASE = Path(__file__).resolve().parent
LAB = BASE.parent
ROOT = LAB.parent
SHOTS = LAB / "screenshots"
REPORT_DIR = LAB / "report"
TEMPLATE = ROOT / "lab-01" / "examples" / "Лаба 2.odt"
FONTS = ROOT / "lab-01" / "build" / "fonts"

sys.path.insert(0, str(BASE))
import content  # noqa: E402

PAGE = {
    "margin_left": 3.0,
    "margin_right": 1.499,
    "margin_top": 2.0,
    "margin_bottom": 2.0,
    "indent": 1.251,
    "font_size": 14,
    "leading": 21.0,
    "image_width": 16.501,
}

STEM = "Отчёт_ЛР2_Основы_работы_в_Linux_GUI_РОСА"
ODT_NAME = STEM + ".odt"
DOCX_NAME = STEM + ".docx"
PDF_NAME = STEM + ".pdf"

# Пространства имён ODT. Их явная регистрация сохраняет привычные префиксы
# в пересобранном content.xml, styles.xml исходного шаблона не изменяется.
NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
    "fo": "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0",
    "svg": "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0",
    "xlink": "http://www.w3.org/1999/xlink",
    "dc": "http://purl.org/dc/elements/1.1/",
    "meta": "urn:oasis:names:tc:opendocument:xmlns:meta:1.0",
    "manifest": "urn:oasis:names:tc:opendocument:xmlns:manifest:1.0",
    "loext": "urn:org:documentfoundation:names:experimental:office:xmlns:loext:1.0",
    "ooo": "http://openoffice.org/2009/office",
}
for _prefix, _url in NS.items():
    ET.register_namespace(_prefix, _url)


def q(prefix: str, local: str) -> str:
    return "{%s}%s" % (NS[prefix], local)


def xml_text(element: ET.Element) -> str:
    return "".join(element.itertext()).strip()


def image_size_cm(name: str, width_cm: float = PAGE["image_width"]) -> tuple[float, float]:
    """Размер снимка при неизменённых пропорциях."""
    with Image.open(SHOTS / name) as im:
        width, height = im.size
    return width_cm, width_cm * height / width


def format_text(text: str, counts: dict[str, int]) -> str:
    return text.format(**counts)


def toc_entries(page_map: dict[str, int | str]) -> list[tuple[int, str, int | str]]:
    """Возвращает строки содержания в порядке документа."""
    result: list[tuple[int, str, int | str]] = []
    for block in content.SECTIONS:
        kind = block[0]
        if kind == "h1":
            result.append((0, block[1], page_map.get(block[1], "")))
        elif kind == "h2":
            result.append((1, block[1], page_map.get(block[1], "")))
        elif kind == "h1c" and len(block) > 2 and block[2]:
            result.append((0, block[1], page_map.get(block[1], "")))
    return result


def count_figures() -> int:
    return sum(1 for block in content.SECTIONS if block[0] == "fig")


def counts_for(pages: int) -> dict[str, int]:
    return {
        "pages": pages,
        "figs": count_figures(),
        "tables": 0,
        "sources": len(content.REFS),
    }


# ------------------------------------------------------------------------------ PDF

def build_pdf(path: Path, counts: dict[str, int]) -> tuple[dict[str, int], int]:
    """Создаёт PDF и возвращает номера страниц заголовков и число страниц."""
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        Image as RLImage,
        KeepTogether,
        NextPageTemplate,
        PageBreak,
        PageTemplate,
        Paragraph,
        Spacer,
    )
    from reportlab.platypus.tableofcontents import TableOfContents

    # Свободный Tinos метрически совместим с Times New Roman; ODT/DOCX при
    # открытии используют именно Times New Roman из стилей шаблона.
    for name, fname in (
        ("Tinos", "Tinos-Regular.ttf"),
        ("Tinos-Bold", "Tinos-Bold.ttf"),
        ("Tinos-Italic", "Tinos-Italic.ttf"),
        ("Tinos-BoldItalic", "Tinos-BoldItalic.ttf"),
    ):
        pdfmetrics.registerFont(TTFont(name, str(FONTS / fname)))
    pdfmetrics.registerFontFamily(
        "Tinos", normal="Tinos", bold="Tinos-Bold", italic="Tinos-Italic", boldItalic="Tinos-BoldItalic"
    )

    size = PAGE["font_size"]
    leading = PAGE["leading"]
    left = PAGE["margin_left"] * cm
    right = PAGE["margin_right"] * cm
    top = PAGE["margin_top"] * cm
    bottom = PAGE["margin_bottom"] * cm

    body = ParagraphStyle(
        "Body",
        fontName="Tinos",
        fontSize=size,
        leading=leading,
        alignment=TA_JUSTIFY,
        firstLineIndent=PAGE["indent"] * cm,
        spaceBefore=0,
        spaceAfter=0,
        splitLongWords=False,
    )
    list_style = ParagraphStyle("List", parent=body)
    h1 = ParagraphStyle(
        "H1",
        parent=body,
        fontName="Tinos-Bold",
        alignment=TA_CENTER,
        firstLineIndent=0,
        keepWithNext=1,
        spaceBefore=0,
        spaceAfter=leading,
    )
    h2 = ParagraphStyle(
        "H2",
        parent=body,
        fontName="Tinos-Bold",
        charSpace=3,
        alignment=TA_JUSTIFY,
        keepWithNext=1,
        spaceBefore=0,
        spaceAfter=leading,
    )
    caption = ParagraphStyle(
        "Caption",
        parent=body,
        alignment=TA_CENTER,
        firstLineIndent=0,
        spaceBefore=0,
        spaceAfter=leading,
    )
    toc1 = ParagraphStyle(
        "TOC1",
        fontName="Tinos",
        fontSize=size,
        leading=leading,
        leftIndent=0,
        firstLineIndent=0,
        rightIndent=1.4 * cm,
    )
    toc2 = ParagraphStyle(
        "TOC2",
        parent=toc1,
        leftIndent=0.423 * cm,
    )
    cover_center = ParagraphStyle(
        "CoverCenter",
        fontName="Tinos",
        fontSize=size,
        leading=leading,
        alignment=TA_CENTER,
        firstLineIndent=0,
        spaceBefore=0,
        spaceAfter=0,
    )
    cover_right = ParagraphStyle("CoverRight", parent=cover_center, alignment=TA_RIGHT)

    class ReportDoc(BaseDocTemplate):
        def __init__(self, filename: str, **kwargs: object) -> None:
            super().__init__(filename, **kwargs)
            self.toc_pages: dict[str, int] = {}

        def afterFlowable(self, flowable: object) -> None:
            entry = getattr(flowable, "_toc_entry", None)
            if entry:
                level, title = entry
                self.toc_pages[title] = self.page
                self.notify("TOCEntry", (level, title, self.page))

    frame = Frame(
        left,
        bottom,
        A4[0] - left - right,
        A4[1] - top - bottom,
        id="main",
        leftPadding=0,
        rightPadding=0,
        topPadding=0,
        bottomPadding=0,
    )

    def draw_page_number(canvas: object, doc: object) -> None:
        canvas.saveState()
        canvas.setFont("Tinos", size)
        canvas.drawCentredString(A4[0] / 2, 1.15 * cm, str(doc.page))
        canvas.restoreState()

    doc = ReportDoc(
        str(path),
        pagesize=A4,
        leftMargin=left,
        rightMargin=right,
        topMargin=top,
        bottomMargin=bottom,
        title="Отчёт по лабораторной работе № 2. Основы работы в Linux с GUI",
        author="Лернер Владислав",
    )
    doc.addPageTemplates(
        [
            PageTemplate(id="cover", frames=[frame], onPage=lambda canvas, doc: None),
            PageTemplate(id="main", frames=[frame], onPage=draw_page_number),
        ]
    )

    story: list[object] = [NextPageTemplate("main")]
    for line in content.TITLE:
        align = line.get("pdf_align", "center")
        style = cover_right if align == "right" else cover_center
        if line.get("pdf_bold") or line.get("pdf_size"):
            style = ParagraphStyle(
                "CoverLine_%d" % len(story),
                parent=style,
                fontName="Tinos-Bold" if line.get("pdf_bold") else "Tinos",
                fontSize=line.get("pdf_size", size),
            )
        text = line["text"] or "&nbsp;"
        story.append(Paragraph(text, style))
        if line.get("gap"):
            story.append(Spacer(1, float(line["gap"])))

    figure_number = 0

    def add_heading(title: str, style: ParagraphStyle, level: int, include: bool = True) -> None:
        paragraph = Paragraph(title, style)
        if include:
            paragraph._toc_entry = (level, title)
        story.append(paragraph)

    def add_figure(filename: str, fig_caption: str) -> None:
        nonlocal figure_number
        figure_number += 1
        width_cm, height_cm = image_size_cm(filename)
        pic = RLImage(str(SHOTS / filename), width=width_cm * cm, height=height_cm * cm)
        cap = Paragraph("Рисунок %d – %s" % (figure_number, fig_caption), caption)
        story.append(KeepTogether([pic, cap]))

    for block in content.SECTIONS:
        kind = block[0]
        if kind == "h1c":
            story.append(PageBreak())
            include = len(block) > 2 and bool(block[2])
            add_heading(block[1], h1, 0, include)
        elif kind == "h1":
            story.append(PageBreak())
            add_heading(block[1], h1, 0)
        elif kind == "h2":
            add_heading(block[1], h2, 1)
        elif kind == "p":
            story.append(Paragraph(format_text(block[1], counts), body))
        elif kind == "list":
            for item in block[1]:
                story.append(Paragraph("– " + item, list_style))
        elif kind == "refs":
            for n, item in enumerate(block[1], 1):
                story.append(Paragraph("%d %s" % (n, item), list_style))
        elif kind == "fig":
            add_figure(block[1], block[2])
        elif kind == "toc":
            toc = TableOfContents()
            toc.dotsMinLevel = 0
            toc.levelStyles = [toc1, toc2]
            story.append(toc)
        else:
            raise ValueError("Неизвестный тип блока: %r" % (kind,))

    doc.multiBuild(story)
    if figure_number != count_figures():
        raise RuntimeError("Число рисунков PDF не совпало с содержанием")
    return doc.toc_pages, doc.page


# ----------------------------------------------------------------------------- DOCX

def build_docx(path: Path, page_map: dict[str, int], counts: dict[str, int]) -> None:
    """Создаёт редактируемую DOCX-копию с параметрами исходного ODT-шаблона."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(PAGE["margin_left"])
    section.right_margin = Cm(PAGE["margin_right"])
    section.top_margin = Cm(PAGE["margin_top"])
    section.bottom_margin = Cm(PAGE["margin_bottom"])
    section.different_first_page_header_footer = True

    def set_font(style: object, bold: bool = False, size_: int = 14) -> None:
        style.font.name = "Times New Roman"
        style.font.size = Pt(size_)
        style.font.bold = bold
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")

    normal = document.styles["Normal"]
    set_font(normal)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(PAGE["indent"])
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def make_style(name: str, **kwargs: object) -> object:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        set_font(style, bool(kwargs.get("bold", False)), int(kwargs.get("size", 14)))
        fmt = style.paragraph_format
        fmt.line_spacing = 1.5
        fmt.space_before = Pt(float(kwargs.get("before", 0)))
        fmt.space_after = Pt(float(kwargs.get("after", 0)))
        fmt.first_line_indent = Cm(float(kwargs.get("indent", PAGE["indent"])))
        fmt.alignment = kwargs.get("align", WD_ALIGN_PARAGRAPH.JUSTIFY)
        fmt.keep_with_next = bool(kwargs.get("keep", False))
        return style

    cover_center = make_style("CoverCenter", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0)
    cover_right = make_style("CoverRight", align=WD_ALIGN_PARAGRAPH.RIGHT, indent=0)
    cover_big = make_style("CoverBig", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0)
    h1 = make_style("MainHeading", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
                    after=PAGE["leading"], keep=True)
    h2 = make_style("SubHeading", bold=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                    after=PAGE["leading"], keep=True)
    caption = make_style("FigureCaption", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
                         after=PAGE["leading"], keep=True)
    toc1 = make_style("TOC1", indent=0)
    toc2 = make_style("TOC2", indent=0.423)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.paragraph_format.line_spacing = 1.0
    page_run = footer.add_run()
    page_run.font.name = "Times New Roman"
    page_run.font.size = Pt(14)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    page_run._r.extend((field_begin, instruction, field_end))

    # Титульный лист.
    for line in content.TITLE:
        align = line.get("pdf_align", "center")
        style = cover_right if align == "right" else cover_center
        if line.get("pdf_bold"):
            style = cover_big
        paragraph = document.add_paragraph(style=style)
        paragraph.add_run(line["text"])
        if line.get("gap"):
            paragraph.paragraph_format.space_after = Pt(float(line["gap"]))

    figure_number = 0

    def heading(title: str, style: object, level: int | None, page_break: bool) -> None:
        paragraph = document.add_paragraph(title, style=style)
        paragraph.paragraph_format.page_break_before = page_break
        if level:
            paragraph.style.paragraph_format.keep_with_next = True

    def figure(filename: str, text: str) -> None:
        nonlocal figure_number
        figure_number += 1
        paragraph = document.add_paragraph(style=caption)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(SHOTS / filename), width=Cm(PAGE["image_width"]))
        caption_par = document.add_paragraph("Рисунок %d – %s" % (figure_number, text), style=caption)
        caption_par.paragraph_format.keep_with_next = False

    for block in content.SECTIONS:
        kind = block[0]
        if kind == "h1c":
            heading(block[1], h1, 1 if len(block) > 2 and block[2] else None, True)
        elif kind == "h1":
            heading(block[1], h1, 1, True)
        elif kind == "h2":
            heading(block[1], h2, 2, False)
        elif kind == "p":
            document.add_paragraph(format_text(block[1], counts), style=normal)
        elif kind == "list":
            for item in block[1]:
                document.add_paragraph("– " + item, style=normal)
        elif kind == "refs":
            for n, item in enumerate(block[1], 1):
                document.add_paragraph("%d %s" % (n, item), style=normal)
        elif kind == "fig":
            figure(block[1], block[2])
        elif kind == "toc":
            for level, title, page in toc_entries(page_map):
                style = toc1 if level == 0 else toc2
                paragraph = document.add_paragraph(style=style)
                paragraph.paragraph_format.tab_stops.add_tab_stop(
                    Cm(16.501 - level * 0.423), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
                )
                paragraph.add_run("%s\t%s" % (title, page))
        else:
            raise ValueError("Неизвестный тип блока: %r" % (kind,))

    if figure_number != count_figures():
        raise RuntimeError("Число рисунков DOCX не совпало с содержанием")
    props = document.core_properties
    props.title = "Отчёт по лабораторной работе № 2. Основы работы в Linux с GUI"
    props.subject = "Основы работы в Linux с GUI. ОС РОСА"
    props.author = "Лернер Владислав"
    props.keywords = "РОСА, Linux, KDE Plasma, VMware Workstation"
    document.save(str(path))


# ------------------------------------------------------------------------------ ODT

def build_odt(path: Path, page_map: dict[str, int], counts: dict[str, int]) -> None:
    """Собирает ODT, полностью сохраняя стили и макеты страниц эталонного ODT."""
    if not TEMPLATE.exists():
        raise FileNotFoundError("Не найден ODT-шаблон: %s" % TEMPLATE)

    with zipfile.ZipFile(TEMPLATE) as source:
        content_root = ET.fromstring(source.read("content.xml"))
        meta_root = ET.fromstring(source.read("meta.xml"))
        manifest_root = ET.fromstring(source.read("META-INF/manifest.xml"))

        document_text = content_root.find(".//office:body/office:text", NS)
        if document_text is None:
            raise RuntimeError("В шаблоне не найдено тело текстового документа")
        template_toc = document_text.find("text:table-of-content", NS)
        if template_toc is None:
            raise RuntimeError("В шаблоне не найдена таблица содержания")
        template_toc = copy.deepcopy(template_toc)

        # Удаляем только прежнее содержание, а не auto-styles: P24, P25, P5,
        # P6, P12 и остальные точные стили шаблона остаются без изменений.
        for child in list(document_text):
            document_text.remove(child)

        auto_styles = content_root.find("office:automatic-styles", NS)
        if auto_styles is None:
            raise RuntimeError("В шаблоне не найдены автоматические стили")

        # Стиль абзаца с рисунком наследует центрирование исходной подписи P12
        # и не позволяет отделить рисунок от следующей подписи переносом страницы.
        # NonTocHeading визуально повторяет P5, но не имеет уровня структуры:
        # «РЕФЕРАТ» и «СОДЕРЖАНИЕ» поэтому не появятся при обновлении оглавления.
        for style_name in ("FigureImage", "NonTocHeading"):
            old_style = auto_styles.find("style:style[@style:name='%s']" % style_name, NS)
            if old_style is not None:
                auto_styles.remove(old_style)
        figure_style = ET.SubElement(
            auto_styles,
            q("style", "style"),
            {
                q("style", "name"): "FigureImage",
                q("style", "family"): "paragraph",
                q("style", "parent-style-name"): "P12",
            },
        )
        ET.SubElement(
            figure_style,
            q("style", "paragraph-properties"),
            {
                q("fo", "keep-with-next"): "always",
                q("fo", "text-align"): "center",
                q("fo", "text-indent"): "0cm",
                q("style", "auto-text-indent"): "false",
            },
        )
        non_toc_heading = ET.SubElement(
            auto_styles,
            q("style", "style"),
            {
                q("style", "name"): "NonTocHeading",
                q("style", "family"): "paragraph",
                q("style", "parent-style-name"): "Standard_20__28_WW_29_",
            },
        )
        ET.SubElement(
            non_toc_heading,
            q("style", "paragraph-properties"),
            {
                q("fo", "line-height"): "150%",
                q("fo", "text-align"): "center",
                q("style", "justify-single-word"): "false",
                q("fo", "keep-together"): "always",
                q("fo", "keep-with-next"): "always",
                q("fo", "break-before"): "page",
                q("fo", "text-indent"): "0cm",
                q("style", "auto-text-indent"): "false",
            },
        )
        ET.SubElement(
            non_toc_heading,
            q("style", "text-properties"),
            {q("fo", "font-weight"): "bold", q("style", "font-weight-asian"): "bold"},
        )

        def add_p(text: str = "", style: str = "P6") -> ET.Element:
            element = ET.SubElement(document_text, q("text", "p"), {q("text", "style-name"): style})
            element.text = text
            return element

        # Имена закладок известны до построения содержания: поэтому ссылки
        # из оглавления работают и до первого автоматического обновления TOC.
        bookmarks = {
            title: "toc_%02d" % number
            for number, (_, title, _) in enumerate(toc_entries(page_map), 1)
        }

        def add_h(text: str, style: str, level: int, make_bookmark: bool) -> ET.Element:
            element = ET.SubElement(
                document_text,
                q("text", "h"),
                {q("text", "style-name"): style, q("text", "outline-level"): str(level)},
            )
            if make_bookmark:
                name = bookmarks[text]
                start = ET.SubElement(element, q("text", "bookmark-start"), {q("text", "name"): name})
                start.tail = text
                ET.SubElement(element, q("text", "bookmark-end"), {q("text", "name"): name})
            else:
                element.text = text
            return element

        def add_title_line(line: dict[str, object]) -> None:
            # Параграфы P24...P31 и текстовые стили T1/T5 взяты из Лаба 2.odt.
            p = add_p(style=str(line["style"]))
            text = str(line["text"])
            if text == "ЛАБОРАТОРНАЯ РАБОТА № 2.":
                s1 = ET.SubElement(p, q("text", "span"), {q("text", "style-name"): "T1"})
                s1.text = "ЛАБОРАТОРНАЯ РАБОТА "
                s2 = ET.SubElement(p, q("text", "span"), {q("text", "style-name"): "T5"})
                s2.text = "№ 2."
            else:
                p.text = text

        figure_number = 0

        def add_figure(filename: str, fig_caption: str) -> None:
            nonlocal figure_number
            figure_number += 1
            width_cm, height_cm = image_size_cm(filename)
            p = add_p(style="FigureImage")
            frame = ET.SubElement(
                p,
                q("draw", "frame"),
                {
                    q("draw", "name"): "Рисунок %d" % figure_number,
                    q("text", "anchor-type"): "as-char",
                    q("svg", "width"): "%.3fcm" % width_cm,
                    q("svg", "height"): "%.3fcm" % height_cm,
                    q("draw", "z-index"): "0",
                },
            )
            ET.SubElement(
                frame,
                q("draw", "image"),
                {
                    q("xlink", "href"): "Pictures/Figure%03d.jpg" % figure_number,
                    q("xlink", "type"): "simple",
                    q("xlink", "show"): "embed",
                    q("xlink", "actuate"): "onLoad",
                },
            )
            add_p("Рисунок %d – %s" % (figure_number, fig_caption), "P12")

        def add_toc() -> None:
            # В исходном файле уже определены шаблоны TOC, стили Contents 1/2
            # и точечный лидер. Сохраняем их, меняем кэшированные строки.
            toc = copy.deepcopy(template_toc)
            source_settings = toc.find("text:table-of-content-source", NS)
            if source_settings is not None:
                source_settings.set(q("text", "use-outline-level"), "true")
                source_settings.set(q("text", "use-index-source-styles"), "false")
            body = toc.find("text:index-body", NS)
            if body is None:
                body = ET.SubElement(toc, q("text", "index-body"))
            else:
                for child in list(body):
                    body.remove(child)
            for level, title, page in toc_entries(page_map):
                p = ET.SubElement(
                    body,
                    q("text", "p"),
                    {q("text", "style-name"): "P3" if level == 0 else "P4"},
                )
                link_attrs = {
                    q("xlink", "type"): "simple",
                    q("xlink", "href"): "#" + bookmarks.get(title, ""),
                    q("text", "style-name"): "Index_20_Link",
                    q("text", "visited-style-name"): "Index_20_Link",
                }
                link = ET.SubElement(p, q("text", "a"), link_attrs)
                link.text = title
                tab = ET.SubElement(link, q("text", "tab"))
                tab.tail = str(page)
            document_text.append(toc)

        # Первая страница остаётся страницей First Page, определённой стилем P24.
        for line in content.TITLE:
            add_title_line(line)

        for block in content.SECTIONS:
            kind = block[0]
            if kind == "h1c":
                include = len(block) > 2 and bool(block[2])
                if include:
                    add_h(block[1], "P5", 1, True)
                else:
                    add_p(block[1], "NonTocHeading")
            elif kind == "h1":
                add_h(block[1], "P5", 1, True)
            elif kind == "h2":
                add_h(block[1], "P20", 2, True)
            elif kind == "p":
                add_p(format_text(block[1], counts), "P6")
            elif kind == "list":
                for item in block[1]:
                    add_p("– " + item, "P6")
            elif kind == "refs":
                for n, item in enumerate(block[1], 1):
                    add_p("%d %s" % (n, item), "P6")
            elif kind == "fig":
                add_figure(block[1], block[2])
            elif kind == "toc":
                add_toc()
            else:
                raise ValueError("Неизвестный тип блока: %r" % (kind,))

        if figure_number != count_figures():
            raise RuntimeError("Число рисунков ODT не совпало с содержанием")

        # Метаданные нового документа.
        title = meta_root.find(".//dc:title", NS)
        if title is not None:
            title.text = "Отчёт по лабораторной работе № 2. Основы работы в Linux с GUI"
        subject = meta_root.find(".//dc:subject", NS)
        if subject is not None:
            subject.text = "ОС РОСА, KDE Plasma, VMware Workstation"
        for tag, value in (("dc:creator", "Лернер Владислав"), ("meta:initial-creator", "Лернер Владислав")):
            element = meta_root.find(".//" + tag, NS)
            if element is not None:
                element.text = value
        # Дата отчёта соответствует дате выполнения, указанной в материалах работы.
        date_text = "2026-09-17T00:00:00Z"
        for tag in ("dc:date", "meta:creation-date"):
            element = meta_root.find(".//" + tag, NS)
            if element is not None:
                element.text = date_text

        # Старые снимки шаблона заменяются ровно 63 подтверждающими скриншотами.
        path_attr = q("manifest", "full-path")
        for entry in list(manifest_root):
            archive_path = entry.get(path_attr, "")
            if archive_path.startswith("Pictures/") or archive_path.startswith("Thumbnails/"):
                manifest_root.remove(entry)
        for number in range(1, figure_number + 1):
            ET.SubElement(
                manifest_root,
                q("manifest", "file-entry"),
                {path_attr: "Pictures/Figure%03d.jpg" % number, q("manifest", "media-type"): "image/jpeg"},
            )

        # Актуальная миниатюра не влияет на печатное оформление, но убирает
        # старую картинку документа-образца в файловом менеджере.
        thumbnail = REPORT_DIR / ".thumbnail_lab2.png"
        REPORT_DIR.mkdir(parents=True, exist_ok=True)
        with Image.open(SHOTS / "1.jpg") as im:
            im.thumbnail((256, 256))
            im.save(thumbnail, "PNG")
        ET.SubElement(
            manifest_root,
            q("manifest", "file-entry"),
            {path_attr: "Thumbnails/thumbnail.png", q("manifest", "media-type"): "image/png"},
        )

        content_bytes = ET.tostring(content_root, encoding="UTF-8", xml_declaration=True)
        meta_bytes = ET.tostring(meta_root, encoding="UTF-8", xml_declaration=True)
        manifest_bytes = ET.tostring(manifest_root, encoding="UTF-8", xml_declaration=True)

        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as output:
            # ODF требует mimetype первым и без сжатия.
            output.writestr("mimetype", source.read("mimetype"), compress_type=zipfile.ZIP_STORED)
            skipped = {"mimetype", "content.xml", "meta.xml", "META-INF/manifest.xml"}
            for info in source.infolist():
                name = info.filename
                if name in skipped or name.startswith("Pictures/") or name.startswith("Thumbnails/"):
                    continue
                output.writestr(info, source.read(name))
            output.writestr("content.xml", content_bytes)
            output.writestr("meta.xml", meta_bytes)
            output.writestr("META-INF/manifest.xml", manifest_bytes)
            for number in range(1, figure_number + 1):
                output.write(SHOTS / ("%d.jpg" % number), "Pictures/Figure%03d.jpg" % number)
            output.write(thumbnail, "Thumbnails/thumbnail.png")

        thumbnail.unlink(missing_ok=True)


# ------------------------------------------------------------------------- orchestration

def validate_inputs() -> None:
    expected = {"%d.jpg" % n for n in range(1, 64)}
    actual = {p.name for p in SHOTS.glob("*.jpg")}
    missing = sorted(expected - actual, key=lambda name: int(name[:-4]))
    if missing:
        raise FileNotFoundError("Не найдены снимки экрана: %s" % ", ".join(missing))
    figures = [block[1] for block in content.SECTIONS if block[0] == "fig"]
    if figures != ["%d.jpg" % n for n in range(1, 64)]:
        raise ValueError("В отчёте должны использоваться снимки 1.jpg ... 63.jpg строго по порядку")


def main() -> None:
    validate_inputs()
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    pdf_path = REPORT_DIR / PDF_NAME
    docx_path = REPORT_DIR / DOCX_NAME
    odt_path = REPORT_DIR / ODT_NAME

    # Номер страниц в реферате уточняется после вёрстки PDF.
    counts = counts_for(0)
    page_map: dict[str, int] = {}
    total_pages = 0
    for _ in range(5):
        page_map, total_pages = build_pdf(pdf_path, counts)
        if counts["pages"] == total_pages:
            break
        counts = counts_for(total_pages)
    else:
        raise RuntimeError("Не удалось стабилизировать число страниц в реферате")

    build_docx(docx_path, page_map, counts)
    build_odt(odt_path, page_map, counts)

    print("PDF : %s" % pdf_path)
    print("DOCX: %s" % docx_path)
    print("ODT : %s" % odt_path)
    print("%d с., %d рис., %d табл., %d источника" % (
        total_pages, counts["figs"], counts["tables"], counts["sources"]
    ))


if __name__ == "__main__":
    main()
